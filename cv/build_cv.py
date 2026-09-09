# -*- coding: utf-8 -*-
"""
Genera el CV estilo Harvard en PDF y Word, en español e inglés.

    python cv/build_cv.py

Sale a public/cv/ (los PDF se sirven desde el sitio). Necesita Playwright
(Chromium) para el PDF y python-docx para el Word; los dos están en el
entorno virtual del backend de SATIUM:
    C:\\Trabajo\\SATIUM\\BACKEND\\SatiumBackend\\.venv\\Scripts\\python.exe
"""
from __future__ import annotations

import asyncio
import html
import os
import sys
from dataclasses import dataclass, field

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "public", "cv")

# ─── Contenido ───────────────────────────────────────────────────────────────


@dataclass
class Job:
    org: str
    role: str
    place: str
    when: str
    bullets: list[str] = field(default_factory=list)


@dataclass
class CV:
    lang: str
    name: str
    contact: list[str]
    summary_title: str
    summary: str
    experience_title: str
    jobs: list[Job]
    projects_title: str
    projects: list[tuple[str, str]]
    education_title: str
    education: list[tuple[str, str, str]]  # (institución, título, cuándo)
    certs_title: str
    certs: str
    skills_title: str
    skills: list[tuple[str, str]]
    languages: str
    interests_title: str
    interests: str
    file: str


CONTACT = [
    "Tepic, Nayarit, México",
    "+52 311 112 0517",
    "axsandovalbe@gmail.com",
    "linkedin.com/in/axel-sandoval-bernal-42b3532a6",
    "axelbernal-dev.web.app",
    "github.com/AxelsBernal",
]

ES = CV(
    lang="es",
    name="Axel Sandoval Bernal",
    contact=CONTACT,
    summary_title="Resumen",
    summary=(
        "Ingeniero en sistemas computacionales y líder técnico full-stack. Dirijo el desarrollo de SATIUM, "
        "un ERP SaaS de cumplimiento fiscal para México, Chile y Perú con catorce aplicaciones en producción "
        "(FastAPI, React, PostgreSQL, Google Cloud). He migrado sistemas heredados cuadrando los datos al centavo "
        "y opero la plataforma con agentes de IA como parte del equipo. Busco un puesto full-stack senior o de "
        "liderazgo técnico, remoto, en México o Estados Unidos."
    ),
    experience_title="Experiencia",
    jobs=[
        Job(
            org="Satium Technologies",
            role="Gerente de Desarrollo Web e IA",
            place="Remoto",
            when="Mar 2026 – actualidad",
            bullets=[
                "Lidero el equipo de ingeniería (3 personas) y agentes de IA; decido arquitectura, prioridades y despliegues de la plataforma.",
                "Construí la API (FastAPI, PostgreSQL en Supabase, Cloud Run) y el frontend (React, TypeScript) de 14 aplicaciones: cuentas por pagar, facturación electrónica, costos, gastos, recursos humanos, visitas y más.",
                "Migré Costos y Gastos de viaje de Toyota desde un sistema Java de quince años: 593,349 filas y $42 M MXN cuadrados al centavo contra el original; 40 tablas, 65 rutas y 24 pantallas.",
                "Llevé el trabajo pesado a Cloud Run Jobs: exportaciones de 16 min a 87 s; regeneré 1.46 M de PDF fiscales; la lista negra del SAT pasó de 5.5 M de renglones a 635 mil vigencias y la consulta de 71 s a 0.5 s.",
                "Implementé facturación electrónica en México (CFDI 4.0 y nómina con PAC), Chile (DTE vía OpenFactura y bóveda del SII) y Perú (SUNAT), y la nómina de ley mexicana completa (ISR, IMSS, IDSE, PTU, finiquitos).",
                "Opero producción: multi-tenant con cinco niveles de acceso, alertas a Slack, respaldos y documentación viva; definí el método de trabajo con agentes de IA (especificación, cuadre, pruebas, verificación en el navegador).",
            ],
        ),
        Job(
            org="SIAREX Technology",
            role="Desarrollador full-stack (prácticas)",
            place="Remoto",
            when="Mar 2025 – Ene 2026",
            bullets=[
                "Backend empresarial en JavaEE con MySQL sobre el sistema que hoy sustituye SATIUM; mantenimiento de bases de datos y control de versiones con GitHub.",
                "Documenté las reglas de negocio del sistema heredado (Costos, Gastos, RH, Visitas) que después migré.",
            ],
        ),
        Job(
            org="Proyectos freelance",
            role="Desarrollador",
            place="Tepic, Nayarit",
            when="2023 – 2024",
            bullets=[
                "Club del Valle: sistema de gestión de socios con React y Firebase, con correos automáticos por Mailjet.",
                "Punto de venta integrado a Google Sheets para La Crudelia Grill; aplicación de escritorio en Java para un consultorio dental.",
            ],
        ),
    ],
    projects_title="Proyectos",
    projects=[
        (
            "Eventleta · eventleta.com · 2026",
            "Plataforma de inscripciones a carreras 5K, 10K y 21K: pagos con Mercado Pago, dorsal con QR y panel del organizador. React, FastAPI, PostgreSQL, Cloud Run. En producción.",
        ),
    ],
    education_title="Educación",
    education=[
        (
            "Instituto Tecnológico de Tepic (TecNM)",
            "Ingeniería en Sistemas Computacionales, especialidad en Desarrollo Web. Titulado en agosto de 2026.",
            "2020 – 2025",
        ),
        (
            "BIG school",
            "Máster en Inteligencia Artificial: agentes y LLM, Claude Code y Codex, automatización con n8n y Make, vibe coding. En curso.",
            "Sep – Dic 2026",
        ),
    ],
    certs_title="Certificaciones",
    certs=(
        "FastAPI: The Complete Course (Udemy, 2025) · Introduction to Data Science (IE University, 2025) · "
        "Diplomado de Metodologías Emprendedoras (TrepCamp, 2025) · Introduction to Cybersecurity (Cisco, 2024)"
    ),
    skills_title="Habilidades",
    skills=[
        ("Backend", "Python, FastAPI, SQLAlchemy, PostgreSQL, Supabase, Pydantic, WeasyPrint, openpyxl"),
        ("Frontend", "TypeScript, React, Vite, Tailwind, shadcn/ui, Three.js, PWA"),
        ("Nube y operación", "Google Cloud Run y Cloud Run Jobs, Cloud Storage, Firebase, Docker, BigQuery, alertas a Slack"),
        ("Fiscal", "CFDI 4.0, SAT, PAC, SII (Chile), SUNAT (Perú), IMSS / IDSE"),
        ("IA", "Claude Code, Codex, OpenAI API, tool calling, Playwright"),
        ("Legado", "Java, JavaEE / Struts, MySQL, Tomcat"),
    ],
    languages="Español e inglés",
    interests_title="Intereses",
    interests="Frontenis de alto rendimiento: seleccionado nacional de Estados Unidos.",
    file="Axel-Bernal-CV-es",
)

EN = CV(
    lang="en",
    name="Axel Sandoval Bernal",
    contact=[
        "Tepic, Nayarit, Mexico",
        "+52 311 112 0517",
        "axsandovalbe@gmail.com",
        "linkedin.com/in/axel-sandoval-bernal-42b3532a6",
        "axelbernal-dev.web.app",
        "github.com/AxelsBernal",
    ],
    summary_title="Summary",
    summary=(
        "Computer systems engineer and full-stack technical lead. I lead the development of SATIUM, a tax-compliance "
        "SaaS ERP for Mexico, Chile and Peru with fourteen applications in production (FastAPI, React, PostgreSQL, "
        "Google Cloud). I have migrated legacy systems reconciling the data to the cent, and I run the platform with "
        "AI agents as part of the team. Looking for a senior full-stack or technical-lead role, remote, in Mexico or the United States."
    ),
    experience_title="Experience",
    jobs=[
        Job(
            org="Satium Technologies",
            role="Web & AI Development Manager",
            place="Remote",
            when="Mar 2026 – present",
            bullets=[
                "Lead the engineering team (3 people) plus AI agents; own architecture, priorities and deployments of the platform.",
                "Built the API (FastAPI, PostgreSQL on Supabase, Cloud Run) and the frontend (React, TypeScript) of 14 applications: accounts payable, e-invoicing, manufacturing costs, travel expenses, human resources, visitor management and more.",
                "Migrated Toyota's Costs and Travel Expenses from a fifteen-year-old Java system: 593,349 rows and MXN $42M reconciled to the cent against the original; 40 tables, 65 routes and 24 screens.",
                "Moved heavy work to Cloud Run Jobs: exports from 16 min to 87 s; regenerated 1.46M tax PDFs; the SAT blacklist went from 5.5M rows to 635K validity periods and lookups from 71 s to 0.5 s.",
                "Implemented e-invoicing in Mexico (CFDI 4.0 and payroll through a PAC), Chile (DTE via OpenFactura and SII vault) and Peru (SUNAT), plus full Mexican statutory payroll (income tax, IMSS, IDSE, profit sharing, severance).",
                "Run production: multi-tenant with five access levels, Slack alerting, backups and living documentation; defined the AI-agent workflow (spec, reconciliation, tests, browser verification).",
            ],
        ),
        Job(
            org="SIAREX Technology",
            role="Full-stack Developer (internship)",
            place="Remote",
            when="Mar 2025 – Jan 2026",
            bullets=[
                "Enterprise backend in JavaEE with MySQL on the system SATIUM now replaces; database maintenance and version control with GitHub.",
                "Documented the legacy business rules (Costs, Expenses, HR, Visits) that I later migrated.",
            ],
        ),
        Job(
            org="Freelance projects",
            role="Developer",
            place="Tepic, Nayarit",
            when="2023 – 2024",
            bullets=[
                "Club del Valle: member management system with React and Firebase, automated emails through Mailjet.",
                "Point of sale integrated with Google Sheets for La Crudelia Grill; Java desktop application for a dental office.",
            ],
        ),
    ],
    projects_title="Projects",
    projects=[
        (
            "Eventleta · eventleta.com · 2026",
            "Registration platform for 5K, 10K and 21K races: Mercado Pago payments, QR bib and organizer dashboard. React, FastAPI, PostgreSQL, Cloud Run. In production.",
        ),
    ],
    education_title="Education",
    education=[
        (
            "Instituto Tecnológico de Tepic (TecNM)",
            "B.Eng. in Computer Systems Engineering, Web Development specialization. Degree awarded August 2026.",
            "2020 – 2025",
        ),
        (
            "BIG school",
            "Master's in Artificial Intelligence: agents and LLMs, Claude Code and Codex, automation with n8n and Make, vibe coding. In progress.",
            "Sep – Dec 2026",
        ),
    ],
    certs_title="Certifications",
    certs=(
        "FastAPI: The Complete Course (Udemy, 2025) · Introduction to Data Science (IE University, 2025) · "
        "Entrepreneurial Methodologies Diploma (TrepCamp, 2025) · Introduction to Cybersecurity (Cisco, 2024)"
    ),
    skills_title="Skills",
    skills=[
        ("Backend", "Python, FastAPI, SQLAlchemy, PostgreSQL, Supabase, Pydantic, WeasyPrint, openpyxl"),
        ("Frontend", "TypeScript, React, Vite, Tailwind, shadcn/ui, Three.js, PWA"),
        ("Cloud & operations", "Google Cloud Run and Cloud Run Jobs, Cloud Storage, Firebase, Docker, BigQuery, Slack alerting"),
        ("Tax", "CFDI 4.0, SAT, PAC, SII (Chile), SUNAT (Peru), IMSS / IDSE"),
        ("AI", "Claude Code, Codex, OpenAI API, tool calling, Playwright"),
        ("Legacy", "Java, JavaEE / Struts, MySQL, Tomcat"),
    ],
    languages="Spanish and English",
    interests_title="Interests",
    interests="High-performance frontenis: United States national team player.",
    file="Axel-Bernal-CV-en",
)

# ─── HTML → PDF ──────────────────────────────────────────────────────────────

CSS = """
@page { size: Letter; margin: 0.42in 0.52in; }
* { box-sizing: border-box; }
body { font-family: Georgia, 'Times New Roman', serif; font-size: 9.5pt; line-height: 1.24; color: #000; margin: 0; }
h1 { font-size: 17.5pt; font-weight: 700; text-align: center; margin: 0 0 2pt; letter-spacing: .02em; }
.contact { text-align: center; font-size: 8.7pt; margin-bottom: 5pt; }
.contact span + span::before { content: '  ·  '; }
h2 { font-size: 9.4pt; font-weight: 700; text-transform: uppercase; letter-spacing: .08em; border-bottom: 1px solid #000; margin: 6.5pt 0 3pt; padding-bottom: 1pt; }
p { margin: 0 0 2pt; text-align: justify; }
.job { margin-bottom: 3.5pt; }
.row { display: flex; justify-content: space-between; align-items: baseline; gap: 12pt; }
.row b { font-size: 9.8pt; }
.row .when { white-space: nowrap; font-size: 9pt; }
.role { font-style: italic; margin-bottom: 1pt; }
ul { margin: 0; padding-left: 13pt; }
li { margin-bottom: 1pt; text-align: justify; }
.skills p { margin-bottom: 1pt; }
.skills b { font-weight: 700; }
"""


def render_html(cv: CV) -> str:
    e = html.escape
    parts = [f"<!doctype html><html lang='{cv.lang}'><head><meta charset='utf-8'><style>{CSS}</style></head><body>"]
    parts.append(f"<h1>{e(cv.name)}</h1>")
    parts.append("<div class='contact'>" + "".join(f"<span>{e(c)}</span>" for c in cv.contact) + "</div>")
    parts.append(f"<h2>{e(cv.summary_title)}</h2><p>{e(cv.summary)}</p>")
    parts.append(f"<h2>{e(cv.experience_title)}</h2>")
    for j in cv.jobs:
        parts.append("<div class='job'>")
        parts.append(f"<div class='row'><b>{e(j.org)}</b><span class='when'>{e(j.when)}</span></div>")
        parts.append(f"<div class='row role'><span>{e(j.role)}</span><span class='when'>{e(j.place)}</span></div>")
        parts.append("<ul>" + "".join(f"<li>{e(b)}</li>" for b in j.bullets) + "</ul></div>")
    parts.append(f"<h2>{e(cv.projects_title)}</h2>")
    for title, desc in cv.projects:
        parts.append(f"<p><b>{e(title)}</b>. {e(desc)}</p>")
    parts.append(f"<h2>{e(cv.education_title)}</h2>")
    for inst, degree, when in cv.education:
        parts.append(f"<div class='job'><div class='row'><b>{e(inst)}</b><span class='when'>{e(when)}</span></div><p>{e(degree)}</p></div>")
    parts.append(f"<h2>{e(cv.certs_title)}</h2><p>{e(cv.certs)}</p>")
    parts.append(f"<h2>{e(cv.skills_title)}</h2><div class='skills'>")
    for k, v in cv.skills:
        parts.append(f"<p><b>{e(k)}:</b> {e(v)}</p>")
    parts.append(f"<p><b>{'Idiomas' if cv.lang == 'es' else 'Languages'}:</b> {e(cv.languages)}</p>")
    parts.append(f"<p><b>{e(cv.interests_title)}:</b> {e(cv.interests)}</p></div>")
    parts.append("</body></html>")
    return "".join(parts)


async def to_pdf(html_text: str, path: str) -> None:
    from playwright.async_api import async_playwright

    async with async_playwright() as p:
        browser = await p.chromium.launch()
        page = await browser.new_page()
        await page.set_content(html_text, wait_until="load")
        await page.pdf(path=path, format="Letter", print_background=True, prefer_css_page_size=True)
        await browser.close()


# ─── Word ────────────────────────────────────────────────────────────────────


def to_docx(cv: CV, path: str) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt

    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Inches(0.55)
        s.left_margin = s.right_margin = Inches(0.6)
    style = doc.styles["Normal"]
    style.font.name = "Georgia"
    style.font.size = Pt(10)
    style.paragraph_format.space_after = Pt(2)

    def para(text="", bold=False, italic=False, size=None, align=None, after=2):
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.bold, r.italic = bold, italic
        if size:
            r.font.size = Pt(size)
        if align is not None:
            p.alignment = align
        p.paragraph_format.space_after = Pt(after)
        return p

    def heading(text):
        p = para(text.upper(), bold=True, size=10, after=2)
        p.paragraph_format.space_before = Pt(8)
        pPr = p._p.get_or_add_pPr()
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        pbdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "000000")):
            bottom.set(qn(k), v)
        pbdr.append(bottom)
        pPr.append(pbdr)

    def two_cols(left, right, bold_left=False, italic=False):
        p = doc.add_paragraph()
        p.paragraph_format.tab_stops.add_tab_stop(Inches(7.3), alignment=2)  # right tab
        r1 = p.add_run(left)
        r1.bold, r1.italic = bold_left, italic
        r2 = p.add_run("\t" + right)
        r2.italic = italic
        r2.font.size = Pt(9.5)
        p.paragraph_format.space_after = Pt(1)

    para(cv.name, bold=True, size=18, align=WD_ALIGN_PARAGRAPH.CENTER, after=1)
    para("  ·  ".join(cv.contact), size=9, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)

    heading(cv.summary_title)
    para(cv.summary)

    heading(cv.experience_title)
    for j in cv.jobs:
        two_cols(j.org, j.when, bold_left=True)
        two_cols(j.role, j.place, italic=True)
        for b in j.bullets:
            bp = doc.add_paragraph(b, style="List Bullet")
            bp.paragraph_format.space_after = Pt(1)
        doc.add_paragraph().paragraph_format.space_after = Pt(0)

    heading(cv.projects_title)
    for title, desc in cv.projects:
        p = doc.add_paragraph()
        p.add_run(title).bold = True
        p.add_run(". " + desc)

    heading(cv.education_title)
    for inst, degree, when in cv.education:
        two_cols(inst, when, bold_left=True)
        para(degree)

    heading(cv.certs_title)
    para(cv.certs)

    heading(cv.skills_title)
    for k, v in cv.skills:
        p = doc.add_paragraph()
        p.add_run(k + ": ").bold = True
        p.add_run(v)
        p.paragraph_format.space_after = Pt(1)
    p = doc.add_paragraph()
    p.add_run(("Idiomas" if cv.lang == "es" else "Languages") + ": ").bold = True
    p.add_run(cv.languages)
    p.paragraph_format.space_after = Pt(1)
    p = doc.add_paragraph()
    p.add_run(cv.interests_title + ": ").bold = True
    p.add_run(cv.interests)
    doc.save(path)


def main() -> None:
    os.makedirs(OUT, exist_ok=True)
    for cv in (ES, EN):
        html_text = render_html(cv)
        pdf = os.path.join(OUT, cv.file + ".pdf")
        asyncio.run(to_pdf(html_text, pdf))
        to_docx(cv, os.path.join(OUT, cv.file + ".docx"))
        print("ok", cv.file, os.path.getsize(pdf) // 1024, "KB")


if __name__ == "__main__":
    sys.exit(main())
